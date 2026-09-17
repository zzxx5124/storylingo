"""從「文字差異.xlsx」讀取 TTS 發音校正對照（原始文字 → 替換文字）。

供 Edge TTS 生成音訊前套用：把原文裡會被唸錯的字，換成符合發音的替換字，
只影響朗讀用文字，不影響顯示的文字。
"""
import os
import threading

from . import settings

_lock = threading.Lock()
_cache: list[tuple[str, str]] | None = None
_cache_mtime: float = -1.0


def _load() -> list[tuple[str, str]]:
    """讀取 xlsx 的「原始文字 / 替換文字」兩欄，回傳排序後的對照表。"""
    global _cache, _cache_mtime
    try:
        import openpyxl
    except ImportError:
        return []

    try:
        mtime = os.path.getmtime(settings.REPLACEMENTS_FILE)
    except OSError:
        return []
    with _lock:
        if _cache is not None and mtime == _cache_mtime:
            return _cache
        pairs = []
        try:
            wb = openpyxl.load_workbook(settings.REPLACEMENTS_FILE, read_only=True, data_only=True)
            ws = wb["replacements"] if "replacements" in wb.sheetnames else wb.active
            for row in ws.iter_rows(values_only=True):
                orig = row[0]
                repl = row[1]
                if orig is None or repl is None:
                    continue
                o = str(orig).strip()
                r = str(repl).strip()
                if not o or not r:
                    continue
                pairs.append((o, r))
            wb.close()
        except Exception:
            return []
        # 長詞優先，避免短詞把長詞的一部分誤替換
        pairs.sort(key=lambda p: len(p[0]), reverse=True)
        _cache = pairs
        _cache_mtime = mtime
        return pairs


def apply(text: str) -> str:
    """依對照表替換發音字。找不到設定檔時回傳原文。"""
    pairs = _load()
    if not pairs or not text:
        return text
    out = text
    for o, r in pairs:
        if o in out:
            out = out.replace(o, r)
    return out


def list_replacements(count: int = 5) -> list[tuple[str, str]]:
    return _load()[:count]


def count() -> int:
    return len(_load())